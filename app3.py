from dotenv import load_dotenv
from openai import OpenAI
import json
import os
import requests
from pypdf import PdfReader
import gradio as gr
from docx import Document

load_dotenv(override=True)

class Me:

    def __init__(self):
        self.openai = OpenAI()
        self.name = "Benjamin Larson"
			
        doc = Document("res/Resume.docx")
        self.resume = "\n".join(
             para.text for para in doc.paragraphs if para.text.strip()
			)
        doc = Document("res/writing.docx")
        self.writing = "\n".join(
			para.text for para in doc.paragraphs if para.text.strip()
			)
    def system_prompt(self):
        system_prompt = f"You are acting as {self.name}. You will be writing cover letters for, \
		job descriptions provided by the prompt. Use {self.name}'s career, background, skills and experience. \
		provided below to help shape a short (less than 1500 character) professional  \
		You are given a sample of {self.name}'s writing style and his resume which you can use to answer questions. \
		Be professional and engaging. don't make up anything not found it the resume. Do not add header, just a simply Dear Hiring Manager "

        system_prompt += f"\n\n##Writingstyle:\n{self.writing}\n\n##Resume:\n{self.resume}\n\n"
       
        return system_prompt
		
    def letter(self,message, history):
        messages = [{"role": "system", "content": self.system_prompt()}] + [{"role": "user", "content": message}]
        
        response = self.openai.chat.completions.create(model="gpt-4o-mini", messages=messages)

        return response.choices[0].message.content
	

if __name__ == "__main__":
    me = Me()
    with gr.Blocks() as demo:
        gr.Markdown(
            """
            # Ben Larson, PhD — Cover Letter Creator
            _Generate tailored, professional cover letters based on your resume and a job description._
            """
        )

        gr.ChatInterface(me.letter)

    demo.launch(server_name="0.0.0.0", server_port=7860)
